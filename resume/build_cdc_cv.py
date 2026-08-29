#!/usr/bin/env python3
"""Build an IIT KGP CDC-style one-page Word CV."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Prabhat_Kumar_Resume.docx"
LOGO = ROOT / "IIT_KGP.png"

NAVY = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x33, 0x33, 0x33)
HEADER_FILL = "B5D4EA"
TABLE_HEADER_FILL = "D6EAF6"
RULE = "7BA7C9"


def set_run(run, *, size=9.0, bold=False, italic=False, color=NAVY, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    run.font.highlight_color = None


def set_paragraph_format(
    p: Paragraph,
    *,
    before=0,
    after=0,
    line=11.2,
    left=0,
    right=0,
    align=WD_ALIGN_PARAGRAPH.LEFT,
):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = Pt(line)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.left_indent = Cm(left)
    pf.right_indent = Cm(right)
    p.alignment = align


def shade_cell(cell, fill: str):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_borders(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        color = kwargs.get("color", "B0C4DE")
        sz = kwargs.get("sz", "4")
        val = kwargs.get("val", "single")
        if kwargs.get("none"):
            val, sz = "nil", "0"
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tc_borders.append(el)
    tc_pr.append(tc_borders)


def set_cell_margins(cell, top=40, bottom=40, left=60, right=60):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.append(mar)


def collapse_table(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tbl_pr.append(borders)
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "0")
    indent.set(qn("w:type"), "dxa")
    tbl_pr.append(indent)
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), "5000")
    width.set(qn("w:type"), "pct")
    tbl_pr.append(width)


def prevent_break(table):
    for row in table.rows:
        tr = row._tr
        tr_pr = tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)


def add_hyperlink(paragraph: Paragraph, text: str, url: str, *, size=8.0, bold=False):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1A365D")
    r_pr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz_cs)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    if bold:
        b = OxmlElement("w:b")
        r_pr.append(b)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_mixed(p: Paragraph, parts, *, size=9.0, color=MUTED):
    for text, bold in parts:
        run = p.add_run(text)
        set_run(run, size=size, bold=bold, color=color)


def section_bar(doc: Document, title: str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    collapse_table(table)
    prevent_break(table)
    cell = table.cell(0, 0)
    shade_cell(cell, HEADER_FILL)
    set_cell_borders(cell, color=RULE, sz="4")
    set_cell_margins(cell, top=18, bottom=18, left=70, right=70)
    p = cell.paragraphs[0]
    set_paragraph_format(p, before=0, after=0, line=12)
    run = p.add_run(title.upper())
    set_run(run, size=9.6, bold=True, color=NAVY, font="Calibri")
    spacer = doc.add_paragraph()
    set_paragraph_format(spacer, before=0, after=0, line=2)


def entry_header(doc: Document, left: str, right: str):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=2.0, after=0.2, line=11.4)
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Cm(18.85), WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(left)
    set_run(run, size=9.15, bold=True, color=NAVY)
    spacer = p.add_run(" ")
    set_run(spacer, size=9.15, bold=True, color=NAVY)
    p.add_run("\t")
    date_run = p.add_run(right)
    set_run(date_run, size=8.8, italic=True, color=MUTED)


def bullet(doc: Document, parts):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_format(p, before=0.15, after=0.15, line=10.85, left=0.32)
    p.paragraph_format.first_line_indent = Cm(-0.30)
    add_mixed(p, parts, size=8.9, color=MUTED)


def skill_line(doc: Document, label: str, body: str):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=0.3, after=0.2, line=10.9)
    lab = p.add_run(label)
    set_run(lab, size=8.9, bold=True, color=NAVY)
    rest = p.add_run(body)
    set_run(rest, size=8.9, bold=False, color=MUTED)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.02)
    section.right_margin = Cm(1.02)
    section.top_margin = Cm(0.52)
    section.bottom_margin = Cm(0.48)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(9.05)
    styles["Normal"].font.color.rgb = MUTED
    bullet_style = styles["List Bullet"]
    bullet_style.font.name = "Calibri"
    bullet_style.font.size = Pt(9.05)
    pPr = bullet_style.element.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:left"), "260")
    ind.set(qn("w:hanging"), "180")

    # ----- Header -----
    header = doc.add_table(rows=1, cols=3)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    collapse_table(header)
    header.columns[0].width = Cm(1.85)
    header.columns[1].width = Cm(10.55)
    header.columns[2].width = Cm(6.50)

    logo_cell = header.cell(0, 0)
    name_cell = header.cell(0, 1)
    contact_cell = header.cell(0, 2)
    for c in (logo_cell, name_cell, contact_cell):
        set_cell_borders(c, none=True)
        set_cell_margins(c, top=0, bottom=0, left=0, right=40)

    logo_p = logo_cell.paragraphs[0]
    set_paragraph_format(logo_p, before=0, after=0, line=14)
    if LOGO.exists():
        logo_p.add_run().add_picture(str(LOGO), width=Inches(0.58))

    name_p = name_cell.paragraphs[0]
    set_paragraph_format(name_p, before=1, after=0, line=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    name_run = name_p.add_run("PRABHAT KUMAR")
    set_run(name_run, size=16.5, bold=True, color=NAVY, font="Calibri")

    deg = name_cell.add_paragraph()
    set_paragraph_format(deg, before=1, after=0, line=12.4, align=WD_ALIGN_PARAGRAPH.CENTER)
    deg_run = deg.add_run("B.Sc. in MATHEMATICS AND COMPUTING")
    set_run(deg_run, size=9.3, bold=True, color=MUTED)

    inst = name_cell.add_paragraph()
    set_paragraph_format(inst, before=0, after=0, line=12.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    inst_run = inst.add_run("Indian Institute of Technology, Kharagpur")
    set_run(inst_run, size=9.0, italic=True, color=MUTED)

    contacts = [
        ("+91-7367905043", None),
        ("sauravsandilya7367905043@gmail.com", "mailto:sauravsandilya7367905043@gmail.com"),
        ("linkedin.com/in/prabhat-kumar-654b7b413", "https://www.linkedin.com/in/prabhat-kumar-654b7b413"),
        ("github.com/Prabhat-190", "https://github.com/Prabhat-190"),
    ]
    for i, (text, url) in enumerate(contacts):
        cp = contact_cell.paragraphs[0] if i == 0 else contact_cell.add_paragraph()
        set_paragraph_format(cp, before=0, after=0, line=11.4, align=WD_ALIGN_PARAGRAPH.RIGHT)
        if url:
            add_hyperlink(cp, text, url, size=8.3)
        else:
            r = cp.add_run(text)
            set_run(r, size=8.3, color=MUTED)

    gap = doc.add_paragraph()
    set_paragraph_format(gap, before=0, after=0, line=4)

    # ----- Education -----
    section_bar(doc, "Education")
    edu = doc.add_table(rows=3, cols=4)
    collapse_table(edu)
    prevent_break(edu)
    widths = [Cm(2.2), Cm(4.6), Cm(8.6), Cm(3.5)]
    for i, w in enumerate(widths):
        edu.columns[i].width = w

    headers = ["Year", "Degree / Exam", "Institute", "CGPA / Marks"]
    rows = [
        ["2027", "B.Sc.", "IIT Kharagpur", "7.29 / 10"],
        ["2023", "AISSCE [CBSE]", "Jawahar Navodaya Vidyalaya, Supaul", "77.8%"],
    ]
    for j, h in enumerate(headers):
        cell = edu.cell(0, j)
        shade_cell(cell, TABLE_HEADER_FILL)
        set_cell_borders(cell, color="9BB8CC", sz="4")
        set_cell_margins(cell, top=20, bottom=20, left=50, right=50)
        p = cell.paragraphs[0]
        set_paragraph_format(p, before=0, after=0, line=11)
        run = p.add_run(h)
        set_run(run, size=8.4, bold=True, color=NAVY)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = edu.cell(i, j)
            set_cell_borders(cell, color="C5D8E6", sz="4")
            set_cell_margins(cell, top=16, bottom=16, left=50, right=50)
            p = cell.paragraphs[0]
            set_paragraph_format(p, before=0, after=0, line=11)
            run = p.add_run(val)
            set_run(run, size=8.5, bold=(j == 0 or j == 3), color=MUTED)

    gap = doc.add_paragraph()
    set_paragraph_format(gap, before=0, after=0, line=3)

    # ----- Internships -----
    section_bar(doc, "Internships")
    entry_header(
        doc,
        "Paycraft  |  Software Engineering Intern  |  ABT-to-NLP Settlement Engine",
        "[ MAY'26 - JUL'26 ]",
    )
    bullet(doc, [
        ("Engineered a ", False),
        ("Java 21 Spring Boot", True),
        (" service that emits ABT-to-NLP files in ", False),
        ("Header-Record-Trailer", True),
        (" format and tracks each file through a ", False),
        ("7-stage", True),
        (" path from generation to callback completion.", False),
    ])
    bullet(doc, [
        ("Designed ", False),
        ("5 REST APIs", True),
        (" and persisted metadata, per-record outcomes, and ", False),
        ("4-digit", True),
        (" daily sequences with ", False),
        ("Spring Data JPA / H2", True),
        (", plus ", False),
        ("SHA-256", True),
        (" checksums.", False),
    ])
    bullet(doc, [
        ("Implemented ", False),
        ("ZIP", True),
        (" packaging, mock ", False),
        ("SFTP", True),
        (" with retry, year-month archival, ", False),
        ("SLF4J", True),
        (" request-id logs, ", False),
        ("OpenAPI", True),
        (", ", False),
        ("JUnit 5", True),
        (", and ", False),
        ("Docker", True),
        (" on ", False),
        ("Railway", True),
        (".", False),
    ])

    # ----- Projects -----
    section_bar(doc, "Projects")

    entry_header(
        doc,
        "Multi-Objective Portfolio Optimisation Engine  |  Bachelor's Thesis",
        "[ JUL'26 - PRESENT ]",
    )
    bullet(doc, [
        ("Built a Python pipeline on a ", False),
        ("10 x 756", True),
        (" return matrix; ", False),
        ("Jarque-Bera", True),
        (" rejected normality for all ", False),
        ("10", True),
        (" assets, so mean-variance understates tail risk.", False),
    ])
    bullet(doc, [
        ("Implemented ", False),
        ("4-objective NSGA-III", True),
        (" in ", False),
        ("pymoo", True),
        (" (max return, min variance, max skewness, min kurtosis) on ", False),
        ("Rp = Rw", True),
        (" in ", False),
        ("O(Tn)", True),
        (", producing ", False),
        ("24", True),
        (" Pareto portfolios.", False),
    ])
    bullet(doc, [
        ("Cut volatility ", False),
        ("22%", True),
        (" vs max-Sharpe (", False),
        ("1.85%", True),
        (" to ", False),
        ("1.45%", True),
        (") and capped the top weight at ", False),
        ("57%", True),
        ("; compared against min-variance and max-Sharpe ", False),
        ("SLSQP", True),
        (" baselines.", False),
    ])

    entry_header(
        doc,
        "Algo Trade Simulator  |  Real-time L2 Execution Cost Engine",
        "[ MAY'25 - AUG'26 ]",
    )
    bullet(doc, [
        ("Built a live simulator on ", False),
        ("Level-2", True),
        (" order-book streams with ", False),
        ("6", True),
        (" cost drivers: slippage, ", False),
        ("Almgren-Chriss", True),
        (" impact, fees, liquidity, volatility, and maker-taker mix.", False),
    ])
    bullet(doc, [
        ("Designed a ", False),
        ("4-layer Redis", True),
        (" Pub/Sub pipeline (ingest, book cache, cost, charts) with exponential-backoff reconnect so ", False),
        ("Dash", True),
        (" still boots if Redis is down.", False),
    ])
    bullet(doc, [
        ("Persisted maker-taker and slippage models with ", False),
        ("joblib", True),
        ("; gated spoofed books when imbalance exceeds ", False),
        ("0.98", True),
        ("; deployed ", False),
        ("Gunicorn / Nginx", True),
        (" with ", False),
        ("100+", True),
        (" tests.", False),
    ])

    entry_header(
        doc,
        "Amazon FinAI  |  Smart Checkout and RAG Assistant",
        "[ JAN'26 - APR'26 ]",
    )
    bullet(doc, [
        ("Built a checkout platform with a ", False),
        ("RAG", True),
        (" assistant over wallet, budget, and policy context using ", False),
        ("ChromaDB", True),
        (" and ", False),
        ("OpenAI", True),
        (".", False),
    ])
    bullet(doc, [
        ("Ranked payment methods on ", False),
        ("3 Redis", True),
        (" signals (gateway success, cashback, gift-card expiry) and extracted bank-offer discounts with ", False),
        ("GPT-4 Vision", True),
        (".", False),
    ])
    bullet(doc, [
        ("Shipped ", False),
        ("3", True),
        (" independently deployable services (", False),
        ("React", True),
        (", ", False),
        ("Node.js", True),
        (", ", False),
        ("FastAPI", True),
        (") with ", False),
        ("Docker Compose", True),
        (".", False),
    ])

    # ----- POR -----
    section_bar(doc, "Positions of Responsibility")
    entry_header(
        doc,
        "Kosi Simanchal Mfg. & Construction Pvt. Ltd.  |  Technical Lead",
        "[ MAY'26 - JUL'26 ]",
    )
    bullet(doc, [
        ("Led architecture and production launch of ", False),
        ("kosisimanchal.com", True),
        (" on ", False),
        ("Next.js 16", True),
        (", ", False),
        ("React 19", True),
        (", and ", False),
        ("Bootstrap 5", True),
        (", shipping ", False),
        ("16+", True),
        (" modules for civil, railway, and IEC trade operations.", False),
    ])
    bullet(doc, [
        ("Hardened ", False),
        ("/api/trade-request", True),
        (" with origin lock, HTML sanitization, ", False),
        ("5 req / 15 min", True),
        (" IP throttle, and ", False),
        ("Nodemailer SMTPS (465)", True),
        (" delivery of B2B inquiries to HQ.", False),
    ])
    bullet(doc, [
        ("Designed a ", False),
        ("4-state", True),
        (" client-portal FSM (login, register, forgot-passkey, sent) with Client / Vendor / Sub-Contractor roles, plus payment, tracking, and careers.", False),
    ])
    bullet(doc, [
        ("Shipped ", False),
        ("GitHub -> Vercel", True),
        (" CI/CD with GoDaddy DNS and SSL; published CIN, GSTIN, IEC, MSME master data. Recognised as Technical Lead (", False),
        ("KSMC/TL/2026/001", True),
        (").", False),
    ])

    # ----- Awards -----
    section_bar(doc, "Awards and Achievements")
    bullet(doc, [
        ("Rated Codeforces ", False),
        ("Expert", True),
        (" with a maximum rating of ", False),
        ("1736", True),
        (" and Global Rank ", False),
        ("245", True),
        (" among ", False),
        ("25,673", True),
        (" in Round 1093; ", False),
        ("500+", True),
        (" DSA problems across LeetCode, Codeforces, CodeChef, and HackerRank.", False),
    ])
    bullet(doc, [
        ("Awarded the ", False),
        ("Reliance Foundation Undergraduate Scholarship", True),
        (" (top ", False),
        ("0.8%", True),
        (" of ", False),
        ("120,000+", True),
        (" applicants); Top ", False),
        ("8%", True),
        (" ", False),
        ("JEE Advanced", True),
        (", Top ", False),
        ("1.2%", True),
        (" ", False),
        ("JEE Main", True),
        (", HackerRank Intern certification.", False),
    ])

    # ----- Skills -----
    section_bar(doc, "Skills and Expertise")
    skill_line(doc, "Languages:  ", "Python  |  Java  |  C++  |  JavaScript  |  SQL")
    skill_line(doc, "LLM, RAG & ML:  ", "OpenAI  |  GPT-4 Vision  |  ChromaDB  |  Scikit-learn  |  Keras  |  Pandas  |  NumPy  |  pymoo / NSGA-III  |  LSTM")
    skill_line(doc, "Backend & Infra:  ", "Spring Boot  |  Spring Data JPA  |  FastAPI  |  Node.js  |  Redis  |  REST  |  Docker  |  Gunicorn  |  Nginx  |  Vercel  |  Railway")
    skill_line(doc, "Frontend & Data:  ", "React.js  |  Next.js  |  Dash  |  Streamlit  |  HTML  |  CSS  |  MySQL  |  MongoDB  |  H2  |  Git  |  Gradle  |  JUnit  |  OpenAPI")

    # ----- Coursework -----
    section_bar(doc, "Coursework Information")
    skill_line(
        doc,
        "Computer Science:  ",
        "Algorithms  |  Data Structures  |  DBMS  |  File Organization  |  Computer Organization  |  Theory of Computation  |  Big Data Analysis  |  OOP  |  OS  |  Networks  |  System Design",
    )
    skill_line(
        doc,
        "Mathematics:  ",
        "Optimization  |  Probability and Statistics  |  Stochastic Processes  |  Linear Algebra  |  Discrete Mathematics  |  Graph Theory",
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
