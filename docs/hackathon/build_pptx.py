"""Build the SkillUp deck. Run from repo root: .venv/bin/python docs/hackathon/build_pptx.py"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "Algo_Trade_Simulator.pptx"
SHOT = HERE / "dashboard.png"

BG = RGBColor(0x04, 0x06, 0x0A)
CARD = RGBColor(0x0D, 0x14, 0x22)
ACCENT = RGBColor(0x00, 0xE5, 0xFF)
GREEN = RGBColor(0x00, 0xE6, 0x76)
RED = RGBColor(0xFF, 0x17, 0x44)
YELLOW = RGBColor(0xFF, 0xD6, 0x00)
PINK = RGBColor(0xFF, 0x33, 0xFF)
TEXT = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xB0, 0xC4, 0xDE)
DIM = RGBColor(0x64, 0x74, 0x8B)

W = Inches(13.333)
H = Inches(7.5)


def _set_run(run, size, color, bold=False):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = "Calibri"


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _bar(slide, color=ACCENT):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.12), H)
    _fill(bar, color)


def _footer(slide, n, total=12):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.1), Inches(10.5), Inches(0.3))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Algo Trade Simulator  ·  SkillUp × IBM SkillsBuild  ·  Prabhat Kumar, IIT Kharagpur"
    _set_run(run, 11, DIM)
    num = slide.shapes.add_textbox(Inches(12.2), Inches(7.1), Inches(0.8), Inches(0.3))
    p = num.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = f"{n}/{total}"
    _set_run(run, 11, DIM)


def _title(slide, text, top=0.28, size=32):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    _set_run(run, size, TEXT, bold=True)
    return box


def _body(slide, lines, left=0.5, top=1.15, width=12.3, height=5.6, size=20, color=MUTED, gap=10):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        run = p.add_run()
        run.text = line
        _set_run(run, size, color)
    return box


def _card(slide, left, top, width, height, title, lines, accent=ACCENT):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height),
    )
    _fill(shape, CARD)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(0.08), Inches(height))
    _fill(stripe, accent)
    box = slide.shapes.add_textbox(
        Inches(left + 0.25), Inches(top + 0.15), Inches(width - 0.4), Inches(height - 0.25),
    )
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    _set_run(run, 16, accent, bold=True)
    for line in lines:
        p = tf.add_paragraph()
        p.space_before = Pt(6)
        run = p.add_run()
        run.text = line
        _set_run(run, 14, MUTED)


def _blank():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    return prs


def _slide(prs):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), W, H)
    _fill(bg, BG)
    # send bg to back
    spTree = slide.shapes._spTree
    sp = bg._element
    spTree.remove(sp)
    spTree.insert(2, sp)
    return slide


def build():
    prs = _blank()

    # 1 title
    s = _slide(prs)
    _bar(s)
    kicker = s.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(12), Inches(0.4))
    p = kicker.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "SKILLUP HACKATHON  ×  IBM SKILLSBUILD"
    _set_run(run, 16, ACCENT, bold=True)
    title = s.shapes.add_textbox(Inches(0.7), Inches(2.3), Inches(12), Inches(1.2))
    p = title.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Algo Trade Simulator"
    _set_run(run, 48, TEXT, bold=True)
    sub = s.shapes.add_textbox(Inches(0.7), Inches(3.5), Inches(12), Inches(0.6))
    p = sub.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Pre-trade cost analysis: slippage, fees, market impact"
    _set_run(run, 22, MUTED)
    _body(s, [
        "Prabhat Kumar  ·  solo  ·  IIT Kharagpur",
        "Does not place trades. Shows what a market order would cost.",
    ], top=4.4, size=18, color=DIM, height=1.2)
    _footer(s, 1)

    # 2 problem
    s = _slide(prs)
    _bar(s, RED)
    _title(s, "The problem")
    _card(s, 0.5, 1.2, 4.0, 5.3, "What people see", [
        "Mid price on a chart or backtest",
        "A fill assumed at that mid",
        "PnL that looks clean on paper",
    ], YELLOW)
    _card(s, 4.7, 1.2, 4.0, 5.3, "What they pay", [
        "Slippage walking the book",
        "Exchange fees by VIP tier",
        "Impact that moves the mid",
    ], RED)
    _card(s, 8.9, 1.2, 4.0, 5.3, "Why it hurts", [
        "A thin book kills a small edge",
        "Live results look worse than the backtest",
        "Hard to tell strategy vs cost",
    ], PINK)
    _footer(s, 2)

    # 3 who
    s = _slide(prs)
    _bar(s)
    _title(s, "Who it is for")
    _body(s, [
        "Students and people writing their first algo — not a broker.",
        "No exchange API keys. Nothing is sent to a venue.",
        "Safe to open in a browser and change size / side / fee tier.",
        "If Redis or a live feed is down, a synthetic book still runs.",
    ], size=22, gap=16)
    _footer(s, 3)

    # 4 solution
    s = _slide(prs)
    _bar(s, GREEN)
    _title(s, "What I built")
    _card(s, 0.5, 1.2, 6.1, 5.3, "Pre-trade cost tool", [
        "Takes a live or synthetic L2 book",
        "Estimates slippage, fees, impact, maker/taker",
        "Shows net cost in dollars and percent",
        "Flags a one-sided or zero-spread book",
        "Saves named setups when Redis is on",
    ], GREEN)
    _card(s, 6.9, 1.2, 5.9, 5.3, "Dashboard", [
        "Ticker: mid, spread, imbalance, depth",
        "Presets: market, scalp, swing, large order",
        "Depth + mid-price + cost charts",
        "History of the last simulations",
        "Health: /healthz and /readyz",
    ], ACCENT)
    _footer(s, 4)

    # 5 architecture
    s = _slide(prs)
    _bar(s)
    _title(s, "How data moves")
    _body(s, [
        "Feed  →  Orderbook  →  TradeSimulator  →  Dash UI",
        "",
        "FeedManager picks Redis, websocket, or a local generator.",
        "TradeSimulator is the only place that runs the four cost models.",
        "MarketState is the lock around the book so the UI and the feed do not race.",
        "create_app() builds one process. Gunicorn imports that factory.",
    ], size=20, gap=12)
    _footer(s, 5)

    # 6 feeds
    s = _slide(prs)
    _bar(s, ACCENT)
    _title(s, "Feeds — why a demo actually loads")
    _card(s, 0.5, 1.2, 3.0, 5.3, "auto", ["Redis if it is set", "else synthetic", "default"], ACCENT)
    _card(s, 3.7, 1.2, 3.0, 5.3, "synthetic", ["Always a local book", "No keys", "Used for judging"], GREEN)
    _card(s, 6.9, 1.2, 3.0, 5.3, "websocket", ["Direct L2 URI", "Reconnects", "Optional"], YELLOW)
    _card(s, 10.1, 1.2, 2.8, 5.3, "redis", ["Subscriber only", "Skipped if unset", "No localhost hang"], PINK)
    _footer(s, 6)

    # 7 models
    s = _slide(prs)
    _bar(s, YELLOW)
    _title(s, "Cost models")
    _card(s, 0.5, 1.2, 6.1, 2.4, "Fees", ["VIP0–VIP5 schedule", "Spot vs futures"], YELLOW)
    _card(s, 6.9, 1.2, 5.9, 2.4, "Slippage", ["Regression if trained", "Else size / vol / imbalance", "Returned as a total $ cost"], ACCENT)
    _card(s, 0.5, 3.85, 6.1, 2.65, "Market impact", ["Almgren–Chriss style", "Temp ~ sqrt(participation)", "Permanent ~ linear"], RED)
    _card(s, 6.9, 3.85, 5.9, 2.65, "Maker / taker", ["Logistic model or heuristic", "Shown as a split, not a fill"], GREEN)
    _footer(s, 7)

    # 8 screenshot
    s = _slide(prs)
    _bar(s)
    _title(s, "Dashboard", size=28)
    if SHOT.exists():
        s.shapes.add_picture(str(SHOT), Inches(0.45), Inches(1.05), width=Inches(12.4))
    _footer(s, 8)

    # 9 sample run
    s = _slide(prs)
    _bar(s, PINK)
    _title(s, "One run — BTC-USDT buy $100, VIP0")
    _card(s, 0.5, 1.2, 4.0, 5.3, "Book", [
        "Mid ~ $64,700",
        "Spread ~ 0.01%",
        "Synthetic L2 feed",
        "Connected, sub-second updates",
    ], ACCENT)
    _card(s, 4.7, 1.2, 4.0, 5.3, "Costs", [
        "Slippage  $0.0062  (0.0062%)",
        "Fees      $0.1201  (0.1201%)",
        "Impact    $0.0055  (0.0055%)",
        "Net       $0.1319  (0.1319%)",
    ], PINK)
    _card(s, 8.9, 1.2, 4.0, 5.3, "What that means", [
        "Fees dominate a small order",
        "Impact stays tiny at $100",
        "Same notional on a thin book would look different",
        "Internal latency 0.75 ms",
    ], GREEN)
    _footer(s, 9)

    # 10 Bob
    s = _slide(prs)
    _bar(s, ACCENT)
    _title(s, "IBM Bob")
    _card(s, 0.5, 1.2, 4.0, 5.3, "Plan", [
        "Map book → net cost",
        "One Railway service",
        "Synthetic if Redis is down",
    ], ACCENT)
    _card(s, 4.7, 1.2, 4.0, 5.3, "Ask", [
        "Slippage was a per-unit offset",
        "Impact was a flat 1%",
        "Kept participation-based AC",
        "Rejected depth-as-ADV",
    ], YELLOW)
    _card(s, 8.9, 1.2, 4.0, 5.3, "Agent", [
        "Feed manager + synthetic book",
        "Skip unconfigured Redis",
        "Local CSS (no CDN)",
        "pytest suite",
    ], GREEN)
    _footer(s, 10)

    # 11 deploy
    s = _slide(prs)
    _bar(s)
    _title(s, "Run and deploy")
    _body(s, [
        "Local:  python -m trading.src.main   →  http://localhost:8050",
        "Docker: docker compose up --build   →  http://localhost:8080",
        "Railway: Dockerfile + railway.toml. Health check is /healthz.",
        "Do not set REDIS_HOST=localhost on Railway.",
        "No API keys. MIT license.",
    ], size=20, gap=14)
    _footer(s, 11)

    # 12 close
    s = _slide(prs)
    _bar(s, GREEN)
    kicker = s.shapes.add_textbox(Inches(0.7), Inches(2.0), Inches(12), Inches(0.4))
    p = kicker.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "THANKS"
    _set_run(run, 16, GREEN, bold=True)
    title = s.shapes.add_textbox(Inches(0.7), Inches(2.5), Inches(12), Inches(1.0))
    p = title.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Algo Trade Simulator"
    _set_run(run, 40, TEXT, bold=True)
    _body(s, [
        "github.com/Prabhat-190/algo-trade-simulator",
        "IBM_BOB.md  ·  docs/hackathon/Algo_Trade_Simulator.pptx",
        "Prabhat Kumar  ·  TradeSim  ·  IIT Kharagpur",
    ], top=3.8, size=20, height=2.2, color=MUTED, gap=12)
    _footer(s, 12)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
